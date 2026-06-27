from openpyxl import load_workbook


workbook = load_workbook(r"C:\Dev\UTM_generator\app\services\users.xlsx")
sheet = workbook.active

for row in sheet.iter_rows(values_only=True):
    print(row)