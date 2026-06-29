from pathlib import Path
import csv

from app.state import AppState


class HashExportService:
    def __init__(self, app_state: AppState):
        self.app_state = app_state

    def export_hashes_to_file(self, output_file_path: str) -> str:
        input_file_path = self.app_state.input_file_path
        hashes = self.app_state.generated_hashes

        if not input_file_path:
            raise ValueError("Путь к входному файлу не указан")

        if not hashes:
            raise ValueError("Список хэшей пустой")

        if not output_file_path:
            raise ValueError("Путь к выходному файлу не указан")

        input_path = Path(input_file_path)
        output_path = Path(output_file_path)

        if not input_path.exists():
            raise FileNotFoundError(f"Входной файл не найден: {input_path}")

        if input_path.suffix.lower() == ".csv":
            if output_path.suffix.lower() != ".csv":
                output_path = output_path.with_suffix(".csv")
            self._export_csv(input_path, output_path, hashes)
            return str(output_path)

        if input_path.suffix.lower() != ".xlsx":
            raise ValueError("Поддерживается только формат .xlsx или .csv")

        if output_path.suffix.lower() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")

        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError as error:
            raise ModuleNotFoundError(
                "Для экспорта Excel-файлов установите зависимость openpyxl"
            ) from error

        workbook = load_workbook(input_path)
        sheet = workbook.active

        header_row = 1
        first_data_row = 2

        email_column_index = None
        for cell in sheet[header_row]:
            if cell.value == "email":
                email_column_index = cell.column
                break

        if email_column_index is None:
            raise ValueError("В файле должен быть столбец email")

        rows_with_email = [
            row_number
            for row_number in range(first_data_row, sheet.max_row + 1)
            if sheet.cell(row=row_number, column=email_column_index).value
        ]

        if len(hashes) != len(rows_with_email):
            raise ValueError(
                f"Количество хэшей не соответствует количеству пользователей в файле. "
                f"Строк пользователей: {len(rows_with_email)}, хэшей: {len(hashes)}"
            )

        hash_column_index = sheet.max_column + 1

        sheet.cell(
            row=header_row,
            column=hash_column_index,
            value="user_hash"
        )

        for row_number, user_hash in zip(rows_with_email, hashes):
            sheet.cell(
                row=row_number,
                column=hash_column_index,
                value=user_hash
            )

        hash_column_letter = sheet.cell(
            row=header_row,
            column=hash_column_index
        ).column_letter

        sheet.column_dimensions[hash_column_letter].width = 80

        workbook.save(output_path)

        return str(output_path)

    def _export_csv(self, input_path: Path, output_path: Path, hashes: list[str]) -> None:
        # CHANGED: CSV upload was already supported, so export now preserves CSV files too.
        with open(input_path, newline="", encoding="utf-8") as input_file:
            reader = csv.DictReader(input_file)

            if not reader.fieldnames:
                raise ValueError("CSV-файл пустой")

            if "email" not in reader.fieldnames:
                raise ValueError("В файле должен быть столбец email")

            rows = list(reader)

        rows_with_email = [row for row in rows if row.get("email")]

        if len(hashes) != len(rows_with_email):
            raise ValueError(
                f"Количество хэшей не соответствует количеству пользователей в файле. "
                f"Строк пользователей: {len(rows_with_email)}, хэшей: {len(hashes)}"
            )

        hash_iterator = iter(hashes)
        for row in rows:
            if row.get("email"):
                row["user_hash"] = next(hash_iterator)

        fieldnames = list(reader.fieldnames)
        if "user_hash" not in fieldnames:
            fieldnames.append("user_hash")

        with open(output_path, "w", newline="", encoding="utf-8") as output_file:
            writer = csv.DictWriter(output_file, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)
