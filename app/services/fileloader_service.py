from pathlib import Path
import csv
from openpyxl import load_workbook


class FileLoaderService:
    def load_users(self, file_path: str) -> list[dict]:

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"Файл не найден: {file_path}")

        if path.suffix.lower() == ".xlsx":
            return self._load_from_xlsx(path)

        if path.suffix.lower() == ".csv":
            return self._load_from_csv(path)

        raise ValueError(f"Неподдерживаемый формат файла: {path.suffix}")

    def _load_from_xlsx(self, path: Path) -> list[dict]:

        workbook = load_workbook(path)
        sheet = workbook.active

        rows = list(sheet.iter_rows(values_only=True))

        if not rows:
            raise ValueError("Файл пустой")

        headers = rows[0]

        if "email" not in headers:
            raise ValueError("В файле должен быть столбец email")

        users = []

        for row in rows[1:]:
            user = dict(zip(headers, row))

            if not user.get("email"):
                continue

            users.append(user)

        return users

    def _load_from_csv(self, path: Path) -> list[dict]:
        users = []
        with open(path, newline="", encoding="utf-8") as csvfile:
            reader = csv.DictReader(csvfile)

            if not reader.fieldnames:
                raise ValueError("CSV-файл пустой")

            if "email" not in reader.fieldnames:
                raise ValueError("В файле должен быть столбец email")

            for row in reader:
                if row.get("email"):
                    users.append(row)

        return users
